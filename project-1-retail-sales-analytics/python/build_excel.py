"""
Builds excel/Retail_Sales_Analysis.xlsx : formula-driven summary tables (SUMIFS), charts,
conditional formatting, a discount what-if model and a Pivot Table how-to.
Run AFTER run_pipeline.py:  python python/build_excel.py
"""
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule, CellIsRule
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parents[1]
f = pd.read_csv(ROOT / "powerbi/data/fact_sales_report.csv", parse_dates=["order_date"])
prod = pd.read_csv(ROOT / "powerbi/data/dim_product.csv")
f = f.merge(prod[["product_id", "unit_price"]], on="product_id")
f["list_revenue"] = (f.quantity * f.unit_price).round(2)

cols = ["order_id", "order_date", "order_year", "order_month", "region", "city", "segment", "category",
        "sub_category", "product_name", "quantity", "discount", "list_revenue", "revenue", "cost", "profit"]
f = f[cols].sort_values(["order_date", "order_id"]).reset_index(drop=True)
n = len(f) + 1  # last row incl. header

FONT, BOLD = "Arial", Font(name="Arial", bold=True)
HDR_FILL, HDR_FONT = PatternFill("solid", fgColor="1F4E79"), Font(name="Arial", bold=True, color="FFFFFF")
INPUT_FONT, INPUT_FILL = Font(name="Arial", color="0000FF", bold=True), PatternFill("solid", fgColor="FFFF00")
thin = Side(style="thin", color="BFBFBF"); BORDER = Border(top=thin, bottom=thin, left=thin, right=thin)

wb = Workbook()

# ======================= README =======================
ws = wb.active; ws.title = "README"
rows = [
    ("Retail Sales Analysis - Excel workbook", None),
    ("Purpose", "Sales, profit and discount analysis for 2022-2024 (synthetic retail data, INR). Built to show filtering, SUMIFS modelling, charts and Pivot Tables."),
    ("Sheets", "Summary = formula-driven KPIs/tables/charts | Discount_WhatIf = scenario model | Sales_Data = cleaned line-item table (tblSales) | Pivot_Guide = how to build the Pivot Tables"),
    ("Data source", "Output of sql/02_data_cleaning.sql (de-duplicated, standardised). Columns revenue, cost, profit, list_revenue are exported data fields; everything on Summary is a live formula."),
    ("Assumption 1", "NULL discounts in the raw data were treated as 0% (35 line items)."),
    ("Assumption 2", "revenue = quantity x unit_price x (1 - discount); cost = quantity x unit_cost; profit = revenue - cost. Taxes, shipping and returns are excluded."),
    ("Legend", "Blue bold on yellow = input you can change (Discount_WhatIf!C4). Black = formula."),
]
for i, (a, b) in enumerate(rows, 1):
    ws.cell(i, 1, a).font = Font(name=FONT, bold=True, size=14 if i == 1 else 11)
    if b:
        ws.cell(i, 2, b).font = Font(name=FONT); ws.cell(i, 2).alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 18; ws.column_dimensions["B"].width = 110
for i in range(2, 8): ws.row_dimensions[i].height = 32

# ======================= DATA =======================
wd = wb.create_sheet("Sales_Data")
wd.append(cols)
for r in f.itertuples(index=False):
    wd.append([r.order_id, r.order_date.to_pydatetime(), int(r.order_year), r.order_month, r.region, r.city, r.segment,
               r.category, r.sub_category, r.product_name, int(r.quantity), float(r.discount),
               float(r.list_revenue), float(r.revenue), float(r.cost), float(r.profit)])
for c in wd[1]: c.font, c.fill = HDR_FONT, HDR_FILL
for row in wd.iter_rows(min_row=2, max_row=n):
    row[1].number_format = "yyyy-mm-dd"; row[11].number_format = "0%"
    for k in (12, 13, 14, 15): row[k].number_format = "#,##0"
for col, w in zip("ABCDEFGHIJKLMNOP", [11, 12, 10, 10, 9, 12, 12, 16, 14, 22, 9, 9, 13, 12, 12, 12]):
    wd.column_dimensions[col].width = w
tab = Table(displayName="tblSales", ref=f"A1:P{n}")
tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
wd.add_table(tab); wd.freeze_panes = "A2"

R = lambda col: f"Sales_Data!${col}$2:${col}${n}"   # absolute range helper
REV, PRO, YEAR, MON, CAT, REG, DIS, LIST, COST = R("N"), R("P"), R("C"), R("D"), R("H"), R("E"), R("L"), R("M"), R("O")

# ======================= SUMMARY =======================
s = wb.create_sheet("Summary", 1)
s["A1"] = "Retail Sales Summary 2022-2024"; s["A1"].font = Font(name=FONT, bold=True, size=16)
def hdr(row, labels, col=1):
    for i, l in enumerate(labels):
        c = s.cell(row, col + i, l); c.font, c.fill, c.border = HDR_FONT, HDR_FILL, BORDER
        c.alignment = Alignment(horizontal="center", wrap_text=True)

# KPIs
hdr(3, ["KPI", "Value"])
kpis = [("Total revenue (Rs)", f"=SUM({REV})", "#,##0"), ("Total profit (Rs)", f"=SUM({PRO})", "#,##0"),
        ("Profit margin", "=B5/B4", "0.0%"), ("Line items sold", f"=COUNTA({R('A')})", "#,##0"),
        ("Average discount", f"=AVERAGE({DIS})", "0.0%"), ("Loss-making line items", f'=COUNTIF({PRO},"<0")', "#,##0"),
        ("% of lines at a loss", "=B9/B7", "0.0%")]
for i, (l, fm, nf) in enumerate(kpis, 4):
    s.cell(i, 1, l).font = BOLD; c = s.cell(i, 2, fm); c.number_format = nf; c.font = Font(name=FONT)
    s.cell(i, 1).border = c.border = BORDER

# Year x Category
hdr(13, ["Category", "2022 revenue", "2023 revenue", "2024 revenue", "Total revenue", "Profit", "Margin", "Revenue share"])
cats = ["Technology", "Furniture", "Office Supplies"]
for i, cat in enumerate(cats, 14):
    s.cell(i, 1, cat).font = BOLD
    for j, y in enumerate((2022, 2023, 2024)):
        s.cell(i, 2 + j, f'=SUMIFS({REV},{CAT},$A{i},{YEAR},{y})').number_format = "#,##0"
    s.cell(i, 5, f"=SUM(B{i}:D{i})").number_format = "#,##0"
    s.cell(i, 6, f'=SUMIFS({PRO},{CAT},$A{i})').number_format = "#,##0"
    s.cell(i, 7, f"=F{i}/E{i}").number_format = "0.0%"
    s.cell(i, 8, f"=E{i}/$E$17").number_format = "0.0%"
s.cell(17, 1, "Total").font = BOLD
for col in "BCDEF": s[f"{col}17"] = f"=SUM({col}14:{col}16)"; s[f"{col}17"].number_format = "#,##0"; s[f"{col}17"].font = BOLD
s["G17"] = "=F17/E17"; s["G17"].number_format = "0.0%"; s["H17"] = "=SUM(H14:H16)"; s["H17"].number_format = "0.0%"
s.cell(18, 1, "YoY growth").font = Font(name=FONT, italic=True)
s["C18"] = "=C17/B17-1"; s["D18"] = "=D17/C17-1"
for c in ("C18", "D18"): s[c].number_format = "0.0%"; s[c].font = Font(name=FONT, italic=True)
s.conditional_formatting.add("G14:G16", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"))

# Region
hdr(21, ["Region", "Revenue", "Profit", "Margin"])
for i, rg in enumerate(["West", "South", "North", "East"], 22):
    s.cell(i, 1, rg).font = BOLD
    s.cell(i, 2, f'=SUMIFS({REV},{REG},$A{i})').number_format = "#,##0"
    s.cell(i, 3, f'=SUMIFS({PRO},{REG},$A{i})').number_format = "#,##0"
    s.cell(i, 4, f"=C{i}/B{i}").number_format = "0.0%"
s.conditional_formatting.add("B22:B25", DataBarRule(start_type="num", start_value=0, end_type="max", color="1F4E79"))

# Discount bands
hdr(28, ["Discount band", "From (>)", "To (<=)", "Revenue", "Profit", "Margin", "Verdict"])
bands = [("No discount", -1, 0), ("1-10%", 0, 0.10), ("11-20%", 0.10, 0.20), ("21%+", 0.20, 1)]
for i, (lab, lo, hi) in enumerate(bands, 29):
    s.cell(i, 1, lab).font = BOLD; s.cell(i, 2, lo).number_format = "0%"; s.cell(i, 3, hi).number_format = "0%"
    s.cell(i, 4, f'=SUMIFS({REV},{DIS},">"&B{i},{DIS},"<="&C{i})').number_format = "#,##0"
    s.cell(i, 5, f'=SUMIFS({PRO},{DIS},">"&B{i},{DIS},"<="&C{i})').number_format = "#,##0"
    s.cell(i, 6, f"=E{i}/D{i}").number_format = "0.0%"
    s.cell(i, 7, f'=IF(F{i}<0,"LOSS-MAKING","Profitable")')
s.conditional_formatting.add("F29:F32", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, color="C00000", bold=True)))
s["A33"] = "Note: band edges are in columns B-C (e.g. 'To' of 10% means discount <= 0.10)."; s["A33"].font = Font(name=FONT, italic=True, size=9)

# Monthly trend
months = sorted(f.order_month.unique())
hdr(3, ["Month", "Revenue", "MoM growth", "3-mo avg"], col=11)
for i, mth in enumerate(months, 4):
    s.cell(i, 11, mth)
    s.cell(i, 12, f'=SUMIFS({REV},{MON},K{i})').number_format = "#,##0"
    if i > 4: s.cell(i, 13, f"=L{i}/L{i-1}-1").number_format = "0.0%"
    if i > 6: s.cell(i, 14, f"=AVERAGE(L{i-2}:L{i})").number_format = "#,##0"
last = 3 + len(months)
s.conditional_formatting.add(f"M5:M{last}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, color="C00000")))

# Charts
ch = BarChart(); ch.type, ch.title, ch.height, ch.width = "col", "Revenue by category and year", 7.5, 14
ch.add_data(Reference(s, min_col=2, max_col=4, min_row=13, max_row=16), titles_from_data=True)
ch.set_categories(Reference(s, min_col=1, min_row=14, max_row=16)); s.add_chart(ch, "A36")
lc = LineChart(); lc.title, lc.height, lc.width = "Monthly revenue with 3-month average", 7.5, 18
lc.add_data(Reference(s, min_col=12, min_row=3, max_row=last), titles_from_data=True)
lc.add_data(Reference(s, min_col=14, min_row=3, max_row=last), titles_from_data=True)
lc.set_categories(Reference(s, min_col=11, min_row=4, max_row=last)); s.add_chart(lc, "K42")
for col, w in zip("ABCDEFGHIJKLMN", [26, 15, 15, 15, 15, 15, 14, 14, 3, 3, 11, 14, 12, 13]): s.column_dimensions[col].width = w
for row in s.iter_rows(min_row=4, max_row=32, max_col=8):
    for c in row:
        if c.font.name != FONT: c.font = Font(name=FONT)
s.freeze_panes = "A3"

# ======================= WHAT-IF =======================
w = wb.create_sheet("Discount_WhatIf", 2)
w["A1"] = "What-if: cap every discount at X%"; w["A1"].font = Font(name=FONT, bold=True, size=14)
w["A2"] = "Change the yellow cell. Assumes unit volumes stay the same, so the result is an UPPER BOUND (real customers may buy less)."; w["A2"].font = Font(name=FONT, italic=True)
w["B4"] = "Discount cap"; w["B4"].font = BOLD
w["C4"] = 0.10; w["C4"].number_format = "0%"; w["C4"].font, w["C4"].fill, w["C4"].border = INPUT_FONT, INPUT_FILL, BORDER
w["C4"].comment = Comment("Input: maximum discount allowed. Try 5%, 10%, 15%, 20%.", "Analyst")
lines = [("Current revenue", f"=SUM({REV})"), ("Current profit", f"=SUM({PRO})"), ("Current margin", "=C7/C6"),
         ("Revenue with cap", f"=SUMPRODUCT({LIST}*(1-{DIS}+({DIS}>C4)*({DIS}-C4)))"),
         ("Profit with cap", f"=C9-SUM({COST})"), ("Margin with cap", "=C10/C9"),
         ("Profit uplift (Rs)", "=C10-C7"), ("Profit uplift (%)", "=C10/C7-1")]
for i, (l, fm) in enumerate(lines, 6):
    w.cell(i, 2, l).font = BOLD; c = w.cell(i, 3, fm); c.font = Font(name=FONT)
    c.number_format = "0.0%" if "margin" in l.lower() or "(%)" in l else "#,##0"
    w.cell(i, 2).border = c.border = BORDER
w.column_dimensions["A"].width = 3; w.column_dimensions["B"].width = 26; w.column_dimensions["C"].width = 18

# ======================= PIVOT GUIDE =======================
g = wb.create_sheet("Pivot_Guide")
steps = [
    "HOW TO BUILD THE PIVOT TABLES (2 minutes each) - Sales_Data is already an Excel Table (tblSales)",
    "",
    "Pivot 1 - Revenue & profit by Category x Year",
    "  1. Click any cell in Sales_Data > Insert > PivotTable > New Worksheet.",
    "  2. Rows: category, sub_category | Columns: order_year | Values: Sum of revenue, Sum of profit.",
    "  3. Right-click a value > Show Values As > '% of Grand Total'. Add Slicers for region and segment (PivotTable Analyze > Insert Slicer).",
    "",
    "Pivot 2 - Discount impact",
    "  1. Rows: discount | Values: Sum of revenue, Sum of profit.",
    "  2. Right-click a discount value > Group... or add a calculated field  Margin = profit / revenue.",
    "  3. Conditional Formatting > Color Scales on the Margin column. Negative margins start at 20% discount.",
    "",
    "Pivot 3 - Monthly trend",
    "  1. Rows: order_month | Values: Sum of revenue | Insert > PivotChart > Line.",
    "",
    "Tip for interviews: cross-check each pivot against the Summary sheet formulas - totals must match exactly.",
]
for i, t in enumerate(steps, 1):
    g.cell(i, 1, t).font = Font(name=FONT, bold=(i == 1 or t.startswith("Pivot")))
g.column_dimensions["A"].width = 140

out = ROOT / "excel/Retail_Sales_Analysis.xlsx"
wb.save(out); print("saved", out, "rows:", n - 1)
